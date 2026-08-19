"""
Извлечение текста из загруженного файла документации закупки
=================================================================

Раньше на вкладке "Разбор документации" единственный способ был —
вручную скопировать и вставить текст. Неудобно и медленно. Эти функции
достают текст сами из загруженного файла (PDF, DOCX, TXT), чтобы
копипастить руками было не обязательно (можно всё равно — как
запасной вариант, поле остаётся редактируемым).

Важное ограничение: работает только с файлами, где текст уже есть в
самом файле (обычный PDF/DOCX, экспортированный из Word/системы
подготовки документации). Если документация — это скан-картинка без
текстового слоя (сфотографированные страницы), извлечение вернёт
пусто — тут нужен уже OCR (как в drawing_anonymizer.py для чертежей),
это отдельная возможная доработка, если понадобится часто.
"""

import io

from docx import Document as DocxDocument
from pypdf import PdfReader


def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages_text = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages_text).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs]
    # существенные условия (обеспечение, сроки) в документации закупки часто
    # оформлены таблицей, а не абзацами — не пропускаем их
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)
    return "\n".join(p for p in parts if p.strip()).strip()


def extract_text_from_upload(filename: str, file_bytes: bytes) -> str:
    """Определяет формат по расширению файла и достаёт текст. Бросает
    понятную ValueError для неподдерживаемых форматов вместо непонятного
    падения глубоко внутри сторонней библиотеки."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    if lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    if lower.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="replace")
    raise ValueError(f"Формат файла не поддерживается: {filename}. Поддерживаются PDF, DOCX, TXT.")


if __name__ == "__main__":
    # 1. DOCX — генерируем реальный .docx через ту же библиотеку, что и
    #    document_generator.py, и проверяем, что текст (включая таблицу) достаётся
    demo_doc = DocxDocument()
    demo_doc.add_paragraph("Извещение о проведении электронного аукциона.")
    demo_doc.add_paragraph("Обеспечение заявки — 45 000 руб.")
    table = demo_doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Срок поставки"
    table.rows[0].cells[1].text = "01.10.2026"
    buf = io.BytesIO()
    demo_doc.save(buf)

    docx_text = extract_text_from_upload("извещение.docx", buf.getvalue())
    print("=== Текст из DOCX ===")
    print(docx_text)
    assert "Обеспечение заявки" in docx_text
    assert "Срок поставки" in docx_text and "01.10.2026" in docx_text, "текст из таблицы тоже должен извлечься"
    print("\nDOCX: проверка пройдена.")

    # 2. PDF — минимальный валидный PDF собран вручную (без внешних PDF-writer
    #    библиотек в этом окружении) с точными смещениями байт в xref-таблице,
    #    текст на английском для надёжности теста самой механики извлечения
    #    (кириллица в реальных PDF заказчиков зависит от того, как их система
    #    его сгенерировала — pypdf поддерживает Unicode для нормально
    #    закодированных PDF, это стандартная и проверенная часть библиотеки,
    #    а не то, что нужно отдельно доказывать самодельным тестом)
    def _build_minimal_pdf(text: str) -> bytes:
        content_stream = f"BT /F1 18 Tf 10 100 Td ({text}) Tj ET".encode()
        objects = [
            b"<</Type/Catalog/Pages 2 0 R>>",
            b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]/Resources<</Font<</F1 4 0 R>>>>/Contents 5 0 R>>",
            b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
            b"<</Length " + str(len(content_stream)).encode() + b">>\nstream\n" + content_stream + b"\nendstream",
        ]
        out = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for i, obj in enumerate(objects, start=1):
            offsets.append(len(out))
            out += f"{i} 0 obj".encode() + obj + b"endobj\n"
        xref_start = len(out)
        out += f"xref\n0 {len(objects) + 1}\n".encode()
        out += b"0000000000 65535 f \n"
        for off in offsets[1:]:
            out += f"{off:010d} 00000 n \n".encode()
        out += b"trailer<</Size " + str(len(objects) + 1).encode() + b"/Root 1 0 R>>\n"
        out += f"startxref\n{xref_start}\n%%EOF".encode()
        return bytes(out)

    minimal_pdf = _build_minimal_pdf("DEMO TENDER DOCUMENT TEXT")
    pdf_text = extract_text_from_upload("извещение.pdf", minimal_pdf)
    print("\n=== Текст из PDF ===")
    print(pdf_text)
    assert "DEMO TENDER DOCUMENT TEXT" in pdf_text
    print("\nPDF: проверка пройдена.")

    # 3. Неподдерживаемый формат — должна быть понятная ошибка
    try:
        extract_text_from_upload("чертёж.dwg", b"whatever")
        raise AssertionError("должна была быть выброшена ValueError")
    except ValueError as e:
        print(f"\nНеподдерживаемый формат: ожидаемая ошибка — {e}")

    print("\nПроверка пройдена: извлечение текста работает для DOCX, PDF, отказ для неизвестных форматов.")
